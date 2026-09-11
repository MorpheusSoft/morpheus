import sys
sys.path.insert(0, '/home/lzambrano/Desarrollo/Morpheus/backend')

import os
from datetime import datetime, timedelta, timezone, date
from decimal import Decimal
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
import openpyxl

from app.main import app
from app.api.deps import SessionLocal
from app.models.core import User, Facility, Supplier, Company
from app.models.purchasing import PurchaseOrder, SellOutAgreement, SellOutAgreementLine
from app.models.inventory import (
    Product, ProductVariant, InventorySnapshot, InventoryAdjustment,
    InventoryAdjustmentLine, AdjustmentReason, Warehouse
)
from app.models.sales import Document, DocumentLine, DocumentType, DocumentState, Customer
from app.models.digital_workers import DigitalWorker, DigitalWorkerActionLog
from app.services.shrinkage_profitability_service import (
    evaluate_variant_shrinkage_and_margin, audit_all_shrinkage_profitability, toggle_purchasing_block
)
from app.services.dead_stock_service import (
    evaluate_variant_dead_stock, audit_all_dead_stock
)
from app.services.sell_out_service import (
    create_sell_out_agreement, settle_sell_out_agreement, conciliate_sell_out_credit_note, get_sell_out_agreements
)
from app.services.monthly_reports_service import (
    generate_monthly_comprehensive_report, execute_scheduled_monthly_job
)
from app.services.mrp_service import MRPService
from app.schemas.sell_out import SellOutAgreementCreate, SellOutAgreementLineCreate, ConciliateCreditNoteRequest
from app.core import security

client = TestClient(app)

@pytest.fixture(scope="module")
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()

@pytest.fixture(scope="module")
def auth_headers(db: Session):
    admin_user = db.query(User).filter(User.email == "admin@morpheus.com").first()
    if not admin_user:
        admin_user = User(
            email="admin@morpheus.com",
            full_name="Administrador Master",
            hashed_password=security.get_password_hash("admin123"),
            is_superuser=True,
            is_active=True
        )
        db.add(admin_user)
        db.commit()

    response = client.post("/api/v1/login/access-token", data={"username": "admin@morpheus.com", "password": "admin123"})
    assert response.status_code == 200
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

@pytest.fixture(scope="module")
def setup_phase4_data(db: Session):
    # Ensure Clara exists
    clara = db.query(DigitalWorker).filter(DigitalWorker.agent_code == "CLARA_COMPRAS").first()
    if not clara:
        admin_user = db.query(User).filter(User.email == "admin@morpheus.com").first()
        clara = DigitalWorker(
            user_id=admin_user.id,
            agent_code="CLARA_COMPRAS",
            display_title="Clara - Compras Estratégicas",
            operational_module="PURCHASING",
            system_prompt="Eres Clara, analista senior de compras y abastecimiento de Neo ERP.",
            is_autonomous_active=True
        )
        db.add(clara)
        db.commit()

    # Facility
    facility = db.query(Facility).filter(Facility.code == "FAC-P4-TEST").first()
    if not facility:
        facility = Facility(
            code="FAC-P4-TEST",
            name="Sede Principal Fase 4",
            address="Av. Bolívar Norte",
            is_distribution_center=False,
            is_active=True
        )
        db.add(facility)
        db.commit()
        db.refresh(facility)

    # Warehouse
    wh = db.query(Warehouse).filter(Warehouse.facility_id == facility.id).first()
    if not wh:
        wh = Warehouse(
            facility_id=facility.id,
            name="Almacén Piso de Venta",
            code="WH-P4"
        )
        db.add(wh)
        db.commit()
        db.refresh(wh)

    # Supplier
    supplier = db.query(Supplier).filter(Supplier.tax_id == "J-44556677-4").first()
    if not supplier:
        supplier = Supplier(
            name="Distribuidora Alimentos Frescos C.A.",
            commercial_name="Alimentos Frescos",
            tax_id="J-44556677-4",
            lead_time_days=3,
            restock_coverage_days=7,
            auto_tune_logistics=True,
            is_active=True
        )
        db.add(supplier)
        db.commit()
        db.refresh(supplier)

    # Product & Variant 1: Perecedero con Merma (Queso Gouda)
    p1 = db.query(Product).filter(Product.name == "Queso Gouda Holandés Test").first()
    if not p1:
        p1 = Product(name="Queso Gouda Holandés Test", is_active=True)
        db.add(p1)
        db.commit()
        db.refresh(p1)

    pv1 = db.query(ProductVariant).filter(ProductVariant.sku == "SKU-GOUDA-P4").first()
    if not pv1:
        pv1 = ProductVariant(
            product_id=p1.id,
            sku="SKU-GOUDA-P4",
            sales_price=Decimal("10.00"),
            replacement_cost=Decimal("8.00"), # Margen bruto teórico = 20%
            is_active=True
        )
        db.add(pv1)
        db.commit()
        db.refresh(pv1)

    # Product & Variant 2: Producto Dead Stock (Vino Reserva Especial)
    p2 = db.query(Product).filter(Product.name == "Vino Tinto Reserva Test").first()
    if not p2:
        p2 = Product(name="Vino Tinto Reserva Test", is_active=True)
        db.add(p2)
        db.commit()
        db.refresh(p2)

    pv2 = db.query(ProductVariant).filter(ProductVariant.sku == "SKU-VINO-DEAD").first()
    if not pv2:
        pv2 = ProductVariant(
            product_id=p2.id,
            sku="SKU-VINO-DEAD",
            sales_price=Decimal("45.00"),
            replacement_cost=Decimal("30.00"),
            is_active=True
        )
        db.add(pv2)
        db.commit()
        db.refresh(pv2)

    # Adjustment Reason: Merma
    reason_merma = db.query(AdjustmentReason).filter(AdjustmentReason.code == "MERMA_TEST").first()
    if not reason_merma:
        reason_merma = AdjustmentReason(
            code="MERMA_TEST",
            name="Merma y Avería Operativa Test",
            default_type="OUT",
            is_active=True
        )
        db.add(reason_merma)
        db.commit()
        db.refresh(reason_merma)

    return {
        "facility": facility,
        "warehouse": wh,
        "supplier": supplier,
        "variant_gouda": pv1,
        "variant_vino": pv2,
        "reason_merma": reason_merma
    }

# ==============================================================================
# TEST 1: Factor de Merma Real en Rentabilidad (Caso 7)
# ==============================================================================
def test_shrinkage_profitability_and_margin_gatekeeper(db: Session, setup_phase4_data):
    """
    Simula un producto con margen bruto del 20% ($10 pvp - $8 costo).
    Se registran 50 unidades vendidas y 20 unidades en merma (Factor Merma = 20/70 = 28.6%).
    Como Merma (28.6%) > Margen Bruto (20%), el Margen Real Neto es NEGATIVO (-8.6%).
    Verifica que Clara bloquea la recompra y emite alerta de destrucción de margen.
    """
    pv = setup_phase4_data["variant_gouda"]
    fac = setup_phase4_data["facility"]
    wh = setup_phase4_data["warehouse"]
    reason = setup_phase4_data["reason_merma"]
    admin = db.query(User).filter(User.email == "admin@morpheus.com").first()

    now = datetime.now(timezone.utc)

    # Limpiar líneas previas para aislar la prueba
    db.query(DocumentLine).filter(DocumentLine.variant_id == pv.id).delete()
    db.query(InventoryAdjustmentLine).filter(InventoryAdjustmentLine.product_variant_id == pv.id).delete()
    db.commit()

    # 1. Registrar ventas de 50 unds
    cust = db.query(Customer).first()
    if not cust:
        cust = Customer(rif="J-00000000-0", name="Cliente Mostrador Test")
        db.add(cust)
        db.commit()

    doc = Document(
        document_number=f"FAC-TEST-MERMA-{int(now.timestamp())}",
        type=DocumentType.INVOICE,
        state=DocumentState.PAID,
        customer_id=cust.id,
        facility_id=fac.id,
        currency_id=1,
        created_at=now - timedelta(days=10)
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    doc_line = DocumentLine(
        document_id=doc.id,
        variant_id=pv.id,
        quantity=Decimal("50.0"),
        unit_price=Decimal("10.00"),
        line_total=Decimal("500.00")
    )
    db.add(doc_line)
    db.commit()

    # 2. Registrar merma de 20 unds en WMS
    adj = InventoryAdjustment(
        number=f"AJ-MERMA-{int(now.timestamp())}",
        facility_id=fac.id,
        warehouse_id=wh.id,
        reason_id=reason.id,
        movement_type="OUT",
        state="APPROVED",
        created_by_id=admin.id,
        approved_by_id=admin.id,
        created_at=now - timedelta(days=5)
    )
    db.add(adj)
    db.commit()
    db.refresh(adj)

    adj_line = InventoryAdjustmentLine(
        adjustment_id=adj.id,
        product_variant_id=pv.id,
        quantity=Decimal("20.0"),
        unit_cost=Decimal("8.00"),
        total_value=Decimal("160.00")
    )
    db.add(adj_line)
    db.commit()

    # 3. Evaluar rentabilidad con el servicio
    res = evaluate_variant_shrinkage_and_margin(db, pv.id, lookback_days=90, auto_block=True)

    assert res["units_sold"] == Decimal("50.0")
    assert res["units_shrinkage"] == Decimal("20.0")
    assert res["gross_margin_pct"] == Decimal("20.0")
    # Factor Merma = (20 / (50 + 20)) * 100 = 28.57%
    assert round(res["shrinkage_pct"], 1) == Decimal("28.6")
    # Margen Real Neto = 20 - 28.57 = -8.57%
    assert res["net_real_margin_pct"] < Decimal("0.0")
    assert res["profitability_status"] == "NEGATIVE_MARGIN"
    assert res["is_blocked_for_purchasing"] is True
    assert "Destrucción de Margen" in res["clara_verdict"]

    db.refresh(pv)
    assert pv.is_blocked_for_purchasing is True
    assert "Margen real negativo" in pv.purchasing_blocked_reason

# ==============================================================================
# TEST 2: Detección de Dead Stock & Bloqueo de Recompra (Caso 4)
# ==============================================================================
def test_dead_stock_detection_and_purchasing_block(db: Session, setup_phase4_data):
    """
    Simula un producto con 15 unidades en almacén pero sin ventas en los últimos 75 días.
    Verifica que Clara lo clasifica como 'DEAD_STOCK', calcula el capital inmovilizado ($450),
    bloquea la recompra y provee recomendación de liquidación/Sell-Out.
    """
    pv = setup_phase4_data["variant_vino"]
    fac = setup_phase4_data["facility"]
    wh = setup_phase4_data["warehouse"]

    # Limpiar ventas previas de este SKU para aislar la prueba
    db.query(DocumentLine).filter(DocumentLine.variant_id == pv.id).delete()
    db.commit()

    # 1. Establecer inventario en mano = 15 unds
    snap = db.query(InventorySnapshot).filter(
        InventorySnapshot.variant_id == pv.id,
        InventorySnapshot.facility_id == fac.id
    ).first()
    if not snap:
        snap = InventorySnapshot(
            variant_id=pv.id,
            facility_id=fac.id,
            stock_qty=Decimal("15.0")
        )
        db.add(snap)
    else:
        snap.stock_qty = Decimal("15.0")
    db.commit()

    # 2. Simular que la última venta fue hace 75 días
    old_date = datetime.now(timezone.utc) - timedelta(days=75)
    cust = db.query(Customer).first()
    doc_old = Document(
        document_number=f"FAC-OLD-DEAD-{int(datetime.now().timestamp())}",
        type=DocumentType.INVOICE,
        state=DocumentState.PAID,
        customer_id=cust.id,
        facility_id=fac.id,
        currency_id=1,
        created_at=old_date
    )
    db.add(doc_old)
    db.commit()
    db.refresh(doc_old)

    d_line = DocumentLine(
        document_id=doc_old.id,
        variant_id=pv.id,
        quantity=Decimal("1.0"),
        unit_price=Decimal("45.00"),
        line_total=Decimal("45.00")
    )
    db.add(d_line)
    db.commit()

    # 3. Evaluar Dead Stock
    item = evaluate_variant_dead_stock(db, pv.id, days_threshold=60, auto_block=True)

    assert item["qty_on_hand"] == Decimal("15.0")
    assert item["days_without_sales"] >= 70
    assert item["dead_stock_status"] == "DEAD_STOCK"
    assert item["stock_valuation_usd"] == Decimal("450.00") # 15 * $30 costo
    assert item["is_blocked_for_purchasing"] is True
    assert "Inmovilizado Crítico" in item["clara_recommended_action"]

    # Probar desbloqueo manual
    unblock_res = toggle_purchasing_block(db, pv.id, is_blocked=False, reason="Aprobado por Gerencia")
    assert unblock_res["is_blocked_for_purchasing"] is False

# ==============================================================================
# TEST 3: Convenios Sell-Out & Conciliación de Notas de Crédito (Caso 9)
# ==============================================================================
def test_sell_out_lifecycle_and_nc_conciliation(db: Session, setup_phase4_data):
    """
    Prueba el ciclo de vida completo de convenios Sell-Out:
    1. Creación del convenio para un rango de 15 días con descuento de $2.00 asumido por el proveedor.
    2. Simulación de 100 unidades vendidas en POS.
    3. Clara liquida el convenio: Reclamo = 100 unds * $2.00 = $200.00.
    4. Conciliación 3-Way con N/C recibida por $200.00 -> MATCH_EXACT.
    """
    supplier = setup_phase4_data["supplier"]
    pv = setup_phase4_data["variant_gouda"]
    fac = setup_phase4_data["facility"]
    cust = db.query(Customer).first()

    today = date.today()
    start_date = today - timedelta(days=14)
    end_date = today

    # 1. Crear Convenio Sell-Out
    payload = SellOutAgreementCreate(
        title="Promoción 15 Días Queso Gouda Sell-Out",
        supplier_id=supplier.id,
        start_date=start_date,
        end_date=end_date,
        lines=[
            SellOutAgreementLineCreate(
                variant_id=pv.id,
                regular_price=Decimal("12.00"),
                promo_price=Decimal("10.00"),
                discount_per_unit=Decimal("2.00"),
                provider_share_pct=Decimal("100.0"),
                provider_share_fixed=Decimal("2.00")
            )
        ]
    )
    agreement = create_sell_out_agreement(db, payload)
    assert agreement.id is not None
    assert agreement.status == "ACTIVE"

    # 2. Registrar ventas durante la vigencia del convenio: 100 unidades
    sale_doc = Document(
        document_number=f"FAC-SELLOUT-PROMO-{int(datetime.now().timestamp())}",
        type=DocumentType.INVOICE,
        state=DocumentState.PAID,
        customer_id=cust.id,
        facility_id=fac.id,
        currency_id=1,
        created_at=datetime.combine(start_date + timedelta(days=5), datetime.min.time()).replace(tzinfo=timezone.utc)
    )
    db.add(sale_doc)
    db.commit()
    db.refresh(sale_doc)

    s_line = DocumentLine(
        document_id=sale_doc.id,
        variant_id=pv.id,
        quantity=Decimal("100.0"),
        unit_price=Decimal("10.00"),
        line_total=Decimal("1000.00")
    )
    db.add(s_line)
    db.commit()

    # 3. Liquidación por Clara
    settle_res = settle_sell_out_agreement(db, agreement.id)
    assert settle_res["ok"] is True
    assert settle_res["total_units_sold"] >= Decimal("100.0")
    assert settle_res["total_claim_amount"] >= Decimal("200.00")

    db.refresh(agreement)
    assert agreement.status == "SETTLED"
    claim_amount = agreement.total_claim_amount

    # 4. Conciliación 3-Way con N/C emitida por el proveedor
    nc_req = ConciliateCreditNoteRequest(
        credit_note_number="NC-PROV-998811",
        credit_note_amount=claim_amount,
        credit_note_date=today,
        notes="Nota de Crédito verificada por Clara"
    )
    concil_res = conciliate_sell_out_credit_note(db, agreement.id, nc_req)
    assert concil_res["ok"] is True
    assert concil_res["conciliation_status"] == "MATCH_EXACT"

    db.refresh(agreement)
    assert agreement.status == "CONCILIATED"
    assert agreement.credit_note_number == "NC-PROV-998811"

# ==============================================================================
# TEST 4: Generador de Reportes Mensuales Programados & a Demanda (Caso 5)
# ==============================================================================
def test_monthly_reports_generator(db: Session, setup_phase4_data):
    """
    Prueba la generación del paquete mensual integral en Excel con sus 5 pestañas:
    1. Resumen ODC, 2. Calibración Proveedores, 3. Mermas y Rentabilidad,
    4. Dead Stock Inmovilizado, 5. Convenios Sell-Out.
    """
    summary = generate_monthly_comprehensive_report(db)

    assert "file_url" in summary
    assert summary["file_url"].endswith(".xlsx")
    assert summary["file_size_kb"] > 0

    # Validar que el archivo físico existe y abrirlo con openpyxl
    file_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
        "backend", "static", "uploads", "exports", summary["file_name"]
    )
    # También verificar ruta absoluta si se guardó en EXPORTS_DIR
    if not os.path.exists(file_path):
        from app.services.monthly_reports_service import EXPORTS_DIR
        file_path = os.path.join(EXPORTS_DIR, summary["file_name"])

    assert os.path.exists(file_path)

    wb = openpyxl.load_workbook(file_path)
    sheet_names = wb.sheetnames
    assert "1. Resumen ODC" in sheet_names
    assert "2. Calibración Proveedores" in sheet_names
    assert "3. Mermas y Rentabilidad" in sheet_names
    assert "4. Dead Stock Inmovilizado" in sheet_names
    assert "5. Convenios Sell-Out" in sheet_names

    # Validar ejecución programada de autómata cron
    cron_res = execute_scheduled_monthly_job(db)
    assert cron_res["file_url"] is not None

# ==============================================================================
# TEST 5: Endpoints REST FastAPI
# ==============================================================================
def test_phase4_api_endpoints(db: Session, setup_phase4_data, auth_headers):
    """Verifica los nuevos endpoints REST de la Fase 4."""
    # 1. GET /api/v1/inventory-intelligence/dead-stock
    resp_dead = client.get("/api/v1/inventory-intelligence/dead-stock", headers=auth_headers)
    assert resp_dead.status_code == 200
    assert "total_dead_stock_items" in resp_dead.json()

    # 2. GET /api/v1/inventory-intelligence/shrinkage-profitability
    resp_shrink = client.get("/api/v1/inventory-intelligence/shrinkage-profitability", headers=auth_headers)
    assert resp_shrink.status_code == 200
    assert "total_skus_evaluated" in resp_shrink.json()

    # 3. GET /api/v1/sell-out/
    resp_so = client.get("/api/v1/sell-out/", headers=auth_headers)
    assert resp_so.status_code == 200
    assert isinstance(resp_so.json(), list)

    # 4. POST /api/v1/inventory-intelligence/generate-monthly-report
    resp_rep = client.post("/api/v1/inventory-intelligence/generate-monthly-report", headers=auth_headers)
    assert resp_rep.status_code == 200
    assert "file_url" in resp_rep.json()

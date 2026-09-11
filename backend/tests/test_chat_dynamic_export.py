import sys
sys.path.insert(0, '/home/lzambrano/Desarrollo/Morpheus/backend')

import os
import pytest
import openpyxl
import xml.etree.ElementTree as ET
from decimal import Decimal
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.api.deps import SessionLocal
from app.models.core import User, Facility, Supplier
from app.models.purchasing import PurchaseOrder, PurchaseOrderLine
from app.models.inventory import Product, ProductVariant
from app.services.dynamic_export_service import (
    generate_excel_export,
    generate_xml_export,
    export_cendi_order_to_excel,
    export_cendi_order_to_xml
)
from app.core import security

client = TestClient(app)

@pytest.fixture(scope="module")
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()

@pytest.fixture(scope="module")
def auth_headers(db: Session):
    admin_user = db.query(User).filter(User.email == "admin@morpheus.com").first()
    if not admin_user:
        admin_user = User(
            email="admin@morpheus.com",
            full_name="Administrador Master",
            hashed_password=security.get_password_hash("admin123"),
            is_superuser=True,
            is_active=True
        )
        db.add(admin_user)
        db.commit()

    response = client.post("/api/v1/login/access-token", data={"username": "admin@morpheus.com", "password": "admin123"})
    assert response.status_code == 200
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

def test_generate_excel_export_structure():
    headers = ["SKU", "Producto", "Cantidad", "Costo ($)", "Subtotal ($)"]
    rows = [
        ["HAR-01", "Harina Pan 1kg", 100, 1.10, 110.00],
        ["ARR-01", "Arroz Blanco 1kg", 50, 0.95, 47.50]
    ]
    res = generate_excel_export(
        title="Reporte de Prueba Neo ERP",
        headers=headers,
        rows=rows,
        filename_prefix="test_excel_gen"
    )

    assert res["format"] == "xlsx"
    assert os.path.exists(res["filepath"])
    assert res["total_records"] == 2
    assert res["download_url"].startswith("/static/uploads/exports/")

    # Read back with openpyxl
    wb = openpyxl.load_workbook(res["filepath"])
    ws = wb.active
    assert "Neo ERP" in str(ws.cell(row=1, column=1).value)
    # Header row is at row 4
    assert ws.cell(row=4, column=1).value == "SKU"
    assert ws.cell(row=4, column=2).value == "Producto"
    # First data row is at row 5
    assert ws.cell(row=5, column=1).value == "HAR-01"

def test_generate_xml_export_structure():
    payload = {
        "Reference": "ODC-2026-TEST",
        "Supplier": {"Name": "Empresas Polar", "RIF": "J-00044767-0"},
        "Items": [
            {"SKU": "HAR-01", "Qty": "100", "Cost": "1.10"},
            {"SKU": "ARR-01", "Qty": "50", "Cost": "0.95"}
        ]
    }
    res = generate_xml_export("PurchaseOrder", payload, filename_prefix="test_xml_gen")

    assert res["format"] == "xml"
    assert os.path.exists(res["filepath"])
    assert res["download_url"].startswith("/static/uploads/exports/")

    # Read back with ElementTree
    tree = ET.parse(res["filepath"])
    root = tree.getroot()
    assert root.tag == "PurchaseOrder"
    assert root.attrib.get("Reference") == "ODC-2026-TEST"

def test_export_cendi_breakdown_excel_and_xml():
    lines = [
        {"sku": "HAR-01", "product_name": "Harina Pan 1kg", "boxes_needed": 10, "total_qty": 100.0, "unit_cost": 1.10, "total_subtotal": 110.0}
    ]
    breakdown = [
        {
            "sku": "HAR-01",
            "product_name": "Harina Pan 1kg",
            "stores": [
                {"facility_name": "Tienda Cumboto", "facility_code": "CAT-01", "qty_needed": 60, "boxes_needed": 6, "subtotal": 66.0, "urgency": "CRITICAL"},
                {"facility_name": "Tienda Trigal", "facility_code": "CAT-11", "qty_needed": 40, "boxes_needed": 4, "subtotal": 44.0, "urgency": "WARNING"}
            ]
        }
    ]

    xl_res = export_cendi_order_to_excel(
        order_reference="ODC-2026-99999",
        supplier_name="Empresas Polar",
        cendi_name="CENDI Principal",
        total_amount=110.0,
        lines=lines,
        distribution_breakdown=breakdown
    )
    assert os.path.exists(xl_res["filepath"])
    wb = openpyxl.load_workbook(xl_res["filepath"])
    assert len(wb.sheetnames) >= 2
    assert "Resumen ODC" in wb.sheetnames
    assert "Desglose por Tienda" in wb.sheetnames

    xml_res = export_cendi_order_to_xml(
        order_reference="ODC-2026-99999",
        supplier_name="Empresas Polar",
        cendi_name="CENDI Principal",
        total_amount=110.0,
        distribution_breakdown=breakdown
    )
    assert os.path.exists(xml_res["filepath"])

def test_api_export_breakdown_endpoint(db: Session, auth_headers):
    import uuid
    # Crear una ODC de prueba con distribution_breakdown
    cendi = db.query(Facility).first()
    sup = db.query(Supplier).first()
    po = PurchaseOrder(
        reference=f"ODC-EXPORT-TEST-{uuid.uuid4().hex[:8].upper()}",
        supplier_id=sup.id if sup else 1,
        dest_facility_id=cendi.id if cendi else 1,
        total_amount=Decimal('500.00'),
        consolidation_mode='CONSOLIDATED_CD',
        distribution_breakdown=[
            {
                "sku": "TEST-SKU",
                "product_name": "Test Product",
                "total_qty": 100,
                "unit_cost": 5.0,
                "total_subtotal": 500.0,
                "stores": [
                    {"facility_name": "Sucursal 1", "facility_code": "S1", "qty_needed": 50, "boxes_needed": 5, "subtotal": 250.0}
                ]
            }
        ]
    )
    db.add(po)
    db.commit()
    db.refresh(po)

    res_xl = client.get(f"/api/v1/purchase-orders/{po.id}/export-breakdown?format=xlsx", headers=auth_headers)
    assert res_xl.status_code == 200
    data_xl = res_xl.json()
    assert data_xl["format"] == "xlsx"
    assert data_xl["download_url"].startswith("/static/uploads/exports/")

    res_xml = client.get(f"/api/v1/purchase-orders/{po.id}/export-breakdown?format=xml", headers=auth_headers)
    assert res_xml.status_code == 200
    data_xml = res_xml.json()
    assert data_xml["format"] == "xml"
    assert data_xml["download_url"].startswith("/static/uploads/exports/")

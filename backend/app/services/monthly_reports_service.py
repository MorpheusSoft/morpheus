import os
from datetime import datetime, timezone, timedelta, date
from decimal import Decimal
from typing import Dict, Any, Optional, List
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.purchasing import PurchaseOrder, SellOutAgreement
from app.models.core import Supplier, Facility, ScheduledReport
from app.models.inventory import ProductVariant, Product, InventorySnapshot
from app.models.digital_workers import DigitalWorker, DigitalWorkerActionLog
from app.services.supplier_logistics_service import audit_supplier_logistics
from app.services.shrinkage_profitability_service import audit_all_shrinkage_profitability
from app.services.dead_stock_service import audit_all_dead_stock

EXPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "static", "uploads", "exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)

# Neo ERP Colors
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
BORDER_THIN = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

def style_sheet_headers(ws, headers: List[str]):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

def auto_fit_columns(ws, min_col=1, max_col=15):
    for col in range(min_col, max_col + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def generate_monthly_comprehensive_report(
    db: Session,
    year: Optional[int] = None,
    month: Optional[int] = None,
    custom_title: Optional[str] = None
) -> Dict[str, Any]:
    """
    Generador del Paquete Mensual Integral de Compras de Neo ERP:
    - Pestaña 1: Resumen ODC e Inversión del Periodo.
    - Pestaña 2: OTIF & Calibración de Reglas Logísticas de Proveedores.
    - Pestaña 3: Auditoría de Mermas y Margen Real Neto.
    - Pestaña 4: Inventario Estancado & Dead Stock.
    - Pestaña 5: Convenios Sell-Out y Liquidación de Notas de Crédito.
    """
    now = datetime.now(timezone.utc)
    target_year = year or (now.year if now.month > 1 else now.year - 1)
    target_month = month or (now.month - 1 if now.month > 1 else 12)

    month_name = datetime(target_year, target_month, 1).strftime("%B %Y")
    report_name = custom_title or f"Informe Ejecutivo de Compras y Rentabilidad - {month_name}"

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    # -------------------------------------------------------------
    # 1. Pestaña: Resumen ODC e Inversión
    # -------------------------------------------------------------
    ws_odc = wb.create_sheet(title="1. Resumen ODC")
    odc_headers = [
        "Nro ODC", "Proveedor", "Sede Destino", "Modo Consolidación", 
        "Estado", "Total USD", "Fecha Emisión", "Recepción / Conciliación"
    ]
    style_sheet_headers(ws_odc, odc_headers)

    # Filtrar órdenes emitidas en el mes
    start_date = datetime(target_year, target_month, 1, tzinfo=timezone.utc)
    if target_month == 12:
        end_date = datetime(target_year + 1, 1, 1, tzinfo=timezone.utc)
    else:
        end_date = datetime(target_year, target_month + 1, 1, tzinfo=timezone.utc)

    orders = db.query(PurchaseOrder).filter(
        PurchaseOrder.created_at >= start_date,
        PurchaseOrder.created_at < end_date
    ).order_by(PurchaseOrder.created_at.desc()).all()

    total_po_amount = Decimal('0.0')
    row_idx = 2
    for o in orders:
        supp_name = o.supplier.name if o.supplier else "N/A"
        fac_name = o.dest_facility.name if o.dest_facility else "N/A"
        amount = Decimal(str(o.total_amount or 0.0))
        total_po_amount += amount
        emitted_str = o.created_at.strftime("%Y-%m-%d") if o.created_at else ""
        concil_str = o.conciliated_at.strftime("%Y-%m-%d") if o.conciliated_at else (o.status if o.status else "")

        ws_odc.append([
            o.reference or f"ODC-{o.id}",
            supp_name,
            fac_name,
            "CENDI CONSOLIDADO" if o.consolidation_mode == 'CONSOLIDATED_CD' else "DIRECTO A TIENDA",
            o.status.upper() if o.status else "DRAFT",
            float(amount),
            emitted_str,
            concil_str
        ])
        # Format currency
        ws_odc.cell(row=row_idx, column=6).number_format = '$#,##0.00'
        if row_idx % 2 == 1:
            for c in range(1, 9):
                ws_odc.cell(row=row_idx, column=c).fill = ZEBRA_FILL
        row_idx += 1

    auto_fit_columns(ws_odc, 1, 8)

    # -------------------------------------------------------------
    # 2. Pestaña: OTIF & Calibración de Proveedores
    # -------------------------------------------------------------
    ws_supp = wb.create_sheet(title="2. Calibración Proveedores")
    supp_headers = [
        "Proveedor", "RIF", "Lead Time Ficha", "Lead Time Real Medido", "Desvío Lead Time",
        "Cadencia Ficha", "Cadencia Real", "Desvío Cadencia", "OTIF Cumplimiento %", "Auto-Calibración Clara"
    ]
    style_sheet_headers(ws_supp, supp_headers)

    active_suppliers = db.query(Supplier).filter(Supplier.is_active == True).limit(20).all()
    row_idx = 2
    for s in active_suppliers:
        audit = audit_supplier_logistics(db, s.id, limit=10)
        ws_supp.append([
            s.name,
            s.tax_id,
            f"{s.lead_time_days or 0} días",
            f"{audit['real_lead_time']} días" if audit['real_lead_time'] is not None else "Sin datos",
            f"{audit['lead_time_deviation']:+d} días" if audit['real_lead_time'] is not None else "0",
            f"{s.restock_coverage_days or 0} días",
            f"{audit['real_restock_days']} días" if audit['real_restock_days'] is not None else "Sin datos",
            f"{audit['restock_deviation']:+d} días" if audit['real_restock_days'] is not None else "0",
            float(audit['on_time_score']),
            "ACTIVADA (Autónomo)" if s.auto_tune_logistics else "ASISTIDA (1-Clic)"
        ])
        ws_supp.cell(row=row_idx, column=9).number_format = '0.0"%"'
        if row_idx % 2 == 1:
            for c in range(1, 11):
                ws_supp.cell(row=row_idx, column=c).fill = ZEBRA_FILL
        row_idx += 1

    auto_fit_columns(ws_supp, 1, 10)

    # -------------------------------------------------------------
    # 3. Pestaña: Auditoría de Mermas & Rentabilidad Real
    # -------------------------------------------------------------
    ws_merma = wb.create_sheet(title="3. Mermas y Rentabilidad")
    merma_headers = [
        "SKU", "Producto", "Precio Venta", "Costo", "Margen Bruto %",
        "Unidades Vendidas", "Unidades Merma", "Pérdida Merma USD", "Factor Merma %",
        "Margen Real Neto %", "Estado Rentabilidad", "Recompra Bloqueada"
    ]
    style_sheet_headers(ws_merma, merma_headers)

    shrinkage_audit = audit_all_shrinkage_profitability(db, lookback_days=90, auto_block=False)
    row_idx = 2
    for item in shrinkage_audit["items"]:
        ws_merma.append([
            item["sku"],
            item["product_name"],
            float(item["sales_price"]),
            float(item["replacement_cost"]),
            float(item["gross_margin_pct"]),
            float(item["units_sold"]),
            float(item["units_shrinkage"]),
            float(item["shrinkage_cost_usd"]),
            float(item["shrinkage_pct"]),
            float(item["net_real_margin_pct"]),
            item["profitability_status"],
            "BLOQUEADO" if item["is_blocked_for_purchasing"] else "HABILITADO"
        ])
        ws_merma.cell(row=row_idx, column=3).number_format = '$#,##0.00'
        ws_merma.cell(row=row_idx, column=4).number_format = '$#,##0.00'
        ws_merma.cell(row=row_idx, column=5).number_format = '0.0"%"'
        ws_merma.cell(row=row_idx, column=8).number_format = '$#,##0.00'
        ws_merma.cell(row=row_idx, column=9).number_format = '0.0"%"'
        ws_merma.cell(row=row_idx, column=10).number_format = '0.0"%"'
        if row_idx % 2 == 1:
            for c in range(1, 13):
                ws_merma.cell(row=row_idx, column=c).fill = ZEBRA_FILL
        row_idx += 1

    auto_fit_columns(ws_merma, 1, 12)

    # -------------------------------------------------------------
    # 4. Pestaña: Inventario Estancado & Dead Stock
    # -------------------------------------------------------------
    ws_dead = wb.create_sheet(title="4. Dead Stock Inmovilizado")
    dead_headers = [
        "SKU", "Producto", "Categoría", "Stock en Mano", "Días Sin Ventas",
        "Última Venta", "Valoración USD", "Estatus", "Recompra Bloqueada", "Recomendación de Clara"
    ]
    style_sheet_headers(ws_dead, dead_headers)

    dead_audit = audit_all_dead_stock(db, days_threshold=60, auto_block=False)
    row_idx = 2
    for d in dead_audit["items"]:
        ws_dead.append([
            d["sku"],
            d["product_name"],
            d["category_name"],
            float(d["qty_on_hand"]),
            d["days_without_sales"],
            d["last_sale_date"].strftime("%Y-%m-%d") if d["last_sale_date"] else "Sin ventas",
            float(d["stock_valuation_usd"]),
            d["dead_stock_status"],
            "BLOQUEADO" if d["is_blocked_for_purchasing"] else "LIBRE",
            d["clara_recommended_action"]
        ])
        ws_dead.cell(row=row_idx, column=7).number_format = '$#,##0.00'
        if row_idx % 2 == 1:
            for c in range(1, 11):
                ws_dead.cell(row=row_idx, column=c).fill = ZEBRA_FILL
        row_idx += 1

    auto_fit_columns(ws_dead, 1, 10)

    # -------------------------------------------------------------
    # 5. Pestaña: Convenios Sell-Out y N/C
    # -------------------------------------------------------------
    ws_so = wb.create_sheet(title="5. Convenios Sell-Out")
    so_headers = [
        "Código", "Título", "Proveedor", "Inicio", "Fin", 
        "Estado", "Total Reclamo USD", "N/C Proveedor", "Monto N/C USD", "Estatus Conciliación"
    ]
    style_sheet_headers(ws_so, so_headers)

    agreements = db.query(SellOutAgreement).order_by(SellOutAgreement.created_at.desc()).all()
    row_idx = 2
    total_sell_out_claim = Decimal('0.0')
    for a in agreements:
        supp_name = a.supplier.name if a.supplier else "N/A"
        claim_amt = Decimal(str(a.total_claim_amount or 0.0))
        nc_amt = Decimal(str(a.credit_note_amount or 0.0))
        total_sell_out_claim += claim_amt

        ws_so.append([
            a.code,
            a.title,
            supp_name,
            a.start_date.strftime("%Y-%m-%d"),
            a.end_date.strftime("%Y-%m-%d"),
            a.status,
            float(claim_amt),
            a.credit_note_number or "PENDIENTE",
            float(nc_amt),
            a.conciliation_status
        ])
        ws_so.cell(row=row_idx, column=7).number_format = '$#,##0.00'
        ws_so.cell(row=row_idx, column=9).number_format = '$#,##0.00'
        if row_idx % 2 == 1:
            for c in range(1, 11):
                ws_so.cell(row=row_idx, column=c).fill = ZEBRA_FILL
        row_idx += 1

    auto_fit_columns(ws_so, 1, 10)

    # Remove the empty first default sheet
    wb.remove(default_sheet)

    # Guardar archivo
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = f"Reporte_Mensual_Clara_Compras_{target_year}_{target_month:02d}_{timestamp}.xlsx"
    file_path = os.path.join(EXPORTS_DIR, file_name)
    wb.save(file_path)

    file_url = f"/static/uploads/exports/{file_name}"
    file_size_kb = round(os.path.getsize(file_path) / 1024, 1)

    summary_info = {
        "title": report_name,
        "month_name": month_name,
        "total_orders": len(orders),
        "total_po_amount_usd": float(total_po_amount),
        "dead_stock_count": dead_audit["total_dead_stock_items"],
        "dead_stock_capital_usd": float(dead_audit["total_capital_immobilized_usd"]),
        "shrinkage_loss_usd": float(shrinkage_audit["total_shrinkage_loss_usd"]),
        "skus_at_margin_risk": shrinkage_audit["skus_at_risk"],
        "sell_out_claim_usd": float(total_sell_out_claim),
        "file_name": file_name,
        "file_url": file_url,
        "file_size_kb": file_size_kb
    }

    return summary_info

def execute_scheduled_monthly_job(db: Session, worker_code: str = "CLARA_COMPRAS") -> Dict[str, Any]:
    """
    Ejecutor del autómata cron mensual para generar el informe y dejar registro en el log de Clara.
    """
    now = datetime.now(timezone.utc)
    summary = generate_monthly_comprehensive_report(db)

    # Actualizar registro en core.scheduled_reports
    sched = db.query(ScheduledReport).filter(ScheduledReport.report_type == "MONTHLY_PURCHASES_AUDIT").first()
    if sched:
        sched.last_generated_at = now
        sched.last_file_path = summary["file_url"]

    # Registrar en DigitalWorkerActionLog
    worker = db.query(DigitalWorker).filter(DigitalWorker.agent_code == worker_code).first()
    action_log = DigitalWorkerActionLog(
        worker_id=worker.id if worker else None,
        action_type="MONTHLY_REPORT_GENERATED",
        target_entity_type="scheduled_report",
        target_entity_id="MONTHLY_PURCHASES_AUDIT",
        severity="INFO",
        summary=f"Clara generó el {summary['title']} ({summary['file_size_kb']} KB). Inversión en ODC: ${summary['total_po_amount_usd']:,.2f}.",
        details=summary,
        recipient_target="GERENCIA_COMPRAS",
        status="COMPLETED"
    )
    db.add(action_log)
    db.commit()

    return summary

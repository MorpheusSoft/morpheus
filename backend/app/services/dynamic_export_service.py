import os
import uuid
from datetime import datetime
from decimal import Decimal
from typing import List, Dict, Any, Optional
import xml.etree.ElementTree as ET
import xml.dom.minidom

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "static", "uploads", "exports"))
os.makedirs(EXPORTS_DIR, exist_ok=True)

def _get_file_size_human(num_bytes: int) -> str:
    for unit in ['B', 'KB', 'MB', 'GB']:
        if abs(num_bytes) < 1024.0:
            return f"{num_bytes:3.1f} {unit}"
        num_bytes /= 1024.0
    return f"{num_bytes:.1f} TB"

def generate_excel_export(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    filename_prefix: str = "export",
    sheet_title: str = "Datos",
    additional_sheets: Optional[List[Dict[str, Any]]] = None
) -> Dict[str, Any]:
    """
    Genera un archivo Excel (.xlsx) con diseño corporativo Neo ERP:
    - Membrete institucional Neo ERP
    - Encabezados azul marino (#1E293B) con tipografía blanca y negrita
    - Bordes delgados y filas alternadas (#F8FAFC)
    - Formateo automático de monedas, números y auto-ajuste de ancho de columnas
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Paleta Neo ERP
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    data_font = Font(name="Segoe UI", size=10, color="1E293B")
    bold_data_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    def style_sheet(target_ws, s_title, s_headers, s_rows):
        # 1. Título Corporativo Neo ERP
        target_ws.append([f"Neo ERP • {s_title}"])
        target_ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(s_headers), 4))
        target_ws.cell(row=1, column=1).font = title_font
        target_ws.cell(row=1, column=1).alignment = Alignment(vertical="center")

        # 2. Subtítulo / Metadata
        now_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        target_ws.append([f"Generado por Asistente Digital Clara • {now_str}"])
        target_ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(s_headers), 4))
        target_ws.cell(row=2, column=1).font = subtitle_font
        target_ws.append([]) # Fila en blanco

        # 3. Encabezados de Columna
        header_row_idx = 4
        target_ws.append(s_headers)
        for col_num in range(1, len(s_headers) + 1):
            cell = target_ws.cell(row=header_row_idx, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        target_ws.row_dimensions[header_row_idx].height = 26

        # 4. Filas de Datos
        for r_idx, row_data in enumerate(s_rows, start=5):
            target_ws.append(row_data)
            target_ws.row_dimensions[r_idx].height = 20
            is_even = (r_idx % 2 == 0)
            for c_idx, val in enumerate(row_data, start=1):
                cell = target_ws.cell(row=r_idx, column=c_idx)
                cell.font = data_font
                cell.border = thin_border
                cell.fill = zebra_fill if is_even else white_fill

                # Formato inteligente según tipo de dato
                if isinstance(val, (float, Decimal)):
                    # Si el encabezado sugiere monto o costo
                    header_name = str(s_headers[c_idx - 1]).lower()
                    if any(k in header_name for k in ["monto", "costo", "precio", "subtotal", "total", "usd", "$"]):
                        cell.number_format = '"$"#,##0.00'
                    else:
                        cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(val, int):
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-ajuste de ancho de columnas
        for col in target_ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 2 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            target_ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Habilitar autofiltro en los datos
        last_col_letter = get_column_letter(len(s_headers))
        target_ws.auto_filter.ref = f"A{header_row_idx}:{last_col_letter}{len(s_rows) + 4}"

    style_sheet(ws, title, headers, rows)

    # Hojas adicionales si existen
    if additional_sheets:
        for extra in additional_sheets:
            extra_ws = wb.create_sheet(title=extra.get("title", "Detalle")[:31])
            style_sheet(
                extra_ws,
                extra.get("title", "Detalle"),
                extra.get("headers", []),
                extra.get("rows", [])
            )

    safe_prefix = "".join(c for c in filename_prefix if c.isalnum() or c in ('_', '-')).strip()
    filename = f"{safe_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:4]}.xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)

    file_size_bytes = os.path.getsize(filepath)
    return {
        "filename": filename,
        "filepath": filepath,
        "download_url": f"/static/uploads/exports/{filename}",
        "format": "xlsx",
        "file_size": _get_file_size_human(file_size_bytes),
        "total_records": len(rows)
    }

def generate_xml_export(
    root_element: str,
    data: Dict[str, Any],
    filename_prefix: str = "export"
) -> Dict[str, Any]:
    """
    Genera un archivo XML estándar B2B con identación y estructura limpia.
    """
    root = ET.Element(root_element)

    def build_tree(element, obj):
        if isinstance(obj, dict):
            for key, val in obj.items():
                clean_key = "".join(c for c in key if c.isalnum() or c == '_').strip()
                if isinstance(val, (dict, list)):
                    sub_el = ET.SubElement(element, clean_key)
                    build_tree(sub_el, val)
                else:
                    # Si es valor primitivo, añadir como atributo si es simple o texto
                    if val is not None:
                        element.set(clean_key, str(val))
        elif isinstance(obj, list):
            for item in obj:
                # Nombre del sub-elemento basado en singular o 'Item'
                sub_name = "Item"
                if element.tag.endswith("s"):
                    sub_name = element.tag[:-1]
                sub_el = ET.SubElement(element, sub_name)
                build_tree(sub_el, item)
        else:
            if obj is not None:
                element.text = str(obj)

    build_tree(root, data)
    xml_bytes = ET.tostring(root, encoding='utf-8')
    dom = xml.dom.minidom.parseString(xml_bytes)
    pretty_xml = dom.toprettyxml(indent="  ", encoding="utf-8")

    safe_prefix = "".join(c for c in filename_prefix if c.isalnum() or c in ('_', '-')).strip()
    filename = f"{safe_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:4]}.xml"
    filepath = os.path.join(EXPORTS_DIR, filename)

    with open(filepath, "wb") as f:
        f.write(pretty_xml)

    file_size_bytes = os.path.getsize(filepath)
    return {
        "filename": filename,
        "filepath": filepath,
        "download_url": f"/static/uploads/exports/{filename}",
        "format": "xml",
        "file_size": _get_file_size_human(file_size_bytes),
        "total_records": 1
    }

def export_cendi_order_to_excel(
    order_reference: str,
    supplier_name: str,
    cendi_name: str,
    total_amount: float,
    lines: List[Dict[str, Any]],
    distribution_breakdown: List[Dict[str, Any]]
) -> Dict[str, Any]:
    """
    Genera el Excel corporativo con dos hojas:
    Hoja 1: Resumen de Orden Consolidada CENDI
    Hoja 2: Matriz de Desglose para Cross-Docking a Tiendas
    """
    # Hoja 1: Renglones maestros
    h1_headers = ["N°", "SKU", "Descripción de Producto", "Empaque", "Bultos Solicitados", "Cantidad Total", "Costo Unitario ($)", "Subtotal ($)"]
    h1_rows = []
    for idx, l in enumerate(lines, start=1):
        h1_rows.append([
            idx,
            l.get("sku", ""),
            l.get("product_name", ""),
            l.get("pack_name", "Und. Base"),
            l.get("boxes_needed", l.get("qty_ordered", 1)),
            float(l.get("total_qty", l.get("expected_base_qty", 0))),
            float(l.get("unit_cost", 0)),
            float(l.get("total_subtotal", l.get("subtotal", 0)))
        ])

    # Hoja 2: Desglose por tienda
    h2_headers = ["SKU", "Producto", "Sucursal / Tienda", "Código Sede", "Cantidad Asignada", "Bultos Asignados", "Subtotal Tienda ($)", "Urgencia"]
    h2_rows = []
    for item in distribution_breakdown:
        sku = item.get("sku", "")
        pname = item.get("product_name", "")
        for s in item.get("stores", []):
            h2_rows.append([
                sku,
                pname,
                s.get("facility_name", ""),
                s.get("facility_code", ""),
                float(s.get("qty_needed", 0)),
                int(s.get("boxes_needed", 0)),
                float(s.get("subtotal", 0)),
                s.get("urgency", "NORMAL")
            ])

    additional_sheets = [
        {
            "title": "Desglose por Tienda",
            "headers": h2_headers,
            "rows": h2_rows
        }
    ]

    title = f"Orden Consolidada CENDI {order_reference} - {supplier_name}"
    return generate_excel_export(
        title=title,
        headers=h1_headers,
        rows=h1_rows,
        filename_prefix=f"ODC_CENDI_{order_reference}",
        sheet_title="Resumen ODC",
        additional_sheets=additional_sheets
    )

def export_cendi_order_to_xml(
    order_reference: str,
    supplier_name: str,
    cendi_name: str,
    total_amount: float,
    distribution_breakdown: List[Dict[str, Any]]
) -> Dict[str, Any]:
    """
    Genera el archivo XML B2B para la ODC Consolidada CENDI.
    """
    payload = {
        "Reference": order_reference,
        "Mode": "CONSOLIDATED_CD",
        "Supplier": {"Name": supplier_name},
        "DestinationCENDI": {"Name": cendi_name},
        "Financials": {"TotalAmountUSD": f"{total_amount:.2f}"},
        "OrderLines": [
            {
                "SKU": item.get("sku"),
                "Product": item.get("product_name"),
                "TotalQuantity": str(item.get("total_qty")),
                "UnitCost": f"{float(item.get('unit_cost', 0)):.4f}",
                "Subtotal": f"{float(item.get('total_subtotal', 0)):.2f}",
                "StoreAllocations": [
                    {
                        "Store": s.get("facility_name"),
                        "Code": s.get("facility_code"),
                        "Quantity": str(s.get("qty_needed")),
                        "Boxes": str(s.get("boxes_needed"))
                    }
                    for s in item.get("stores", [])
                ]
            }
            for item in distribution_breakdown
        ]
    }

    return generate_xml_export("PurchaseOrder", payload, filename_prefix=f"B2B_ODC_{order_reference}")
